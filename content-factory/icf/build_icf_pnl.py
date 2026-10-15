"""Build the ICF/IID all-states Pro Forma P&L (state dropdown, resident count, 3-year outlook).

python3 build_icf_pnl.py OUTDIR            -> master file + one preset copy per state slug
"""
import sys, os, re, subprocess, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(__file__))
from icf_rates import R, W, DEFAULT_RATE

NAVY = "1F3A5F"; BLUE_IN = "0000CC"; GREEN = "006100"
F_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
F_SUB = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
F_SEC = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_B = Font(name="Calibri", size=10, bold=True)
F_N = Font(name="Calibri", size=10)
F_IN = Font(name="Calibri", size=10, color=BLUE_IN, bold=True)
F_LK = Font(name="Calibri", size=10, color=GREEN)
F_NOTE = Font(name="Calibri", size=9, italic=True, color="555555")
FILL_T = PatternFill("solid", fgColor=NAVY)
FILL_S = PatternFill("solid", fgColor="2E5C8A")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_TOT = PatternFill("solid", fgColor="DDE7F2")
FILL_HD = PatternFill("solid", fgColor="EAF0F7")
THIN = Side(style="thin", color="B7C4D4")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
USD = '"$"#,##0;[Red]("$"#,##0)'
USD2 = '"$"#,##0.00'
PCT = '0.0%'

MONTHS = [f"M{i}" for i in range(1, 13)]
MC = [L(c) for c in range(2, 14)]      # B..M
Y1, Y2, Y3 = "N", "O", "P"
ALLC = MC + [Y1, Y2, Y3]

STATES = sorted(set(R) | set(W))


def build(state_default, path):
    wb = Workbook()
    ws = wb.active; ws.title = "P&L"
    rl = wb.create_sheet("ICF Rate Lookup")
    wl = wb.create_sheet("Wage Lookup")
    ins = wb.create_sheet("Instructions")

    # ---------- lookup sheets ----------
    rl["A1"] = "ICF/IID Small-Facility Medicaid Daily Rates by State"; rl["A1"].font = Font(size=14, bold=True, color=NAVY)
    rl["A2"] = (f"Researched October 2026 from state rate schedules, rate letters and state notices. 'Default' rows have no usable "
                f"published figure and use a ${DEFAULT_RATE:,.0f}/day planning figure - enter your facility's actual rate on the P&L.")
    rl["A2"].font = F_NOTE
    hdr = ["State", "Daily Rate Used", "Status", "Published / Proxy Figure", "Basis", "Effective", "Small ICF Market",
           "Source", "Source Link", "Notes"]
    for i, h in enumerate(hdr, 1):
        c = rl.cell(4, i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_S; c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, st in enumerate(STATES, 5):
        rate, status, basis, eff, stitle, surl, note, market = R.get(st, (None, "default", "", "", "Not researched", "", "", ""))
        used = rate if rate is not None else DEFAULT_RATE
        label = {"published": "Published", "proxy": "Proxy - low confidence", "floor": "Statutory floor",
                 "default": "Planning default - verify"}[status]
        vals = [st, used, label, rate if rate is not None else "n/a", basis, eff,
                "No small private ICF/IID program found" if market == "none" else "", stitle, surl, note]
        for i, v in enumerate(vals, 1):
            c = rl.cell(r, i, v); c.font = F_N; c.alignment = Alignment(wrap_text=i in (5, 8, 10), vertical="top")
        rl.cell(r, 2).number_format = USD2
        if isinstance(rl.cell(r, 4).value, (int, float)):
            rl.cell(r, 4).number_format = USD2
        if surl:
            rl.cell(r, 9).hyperlink = surl; rl.cell(r, 9).font = Font(size=10, color="0563C1", underline="single")
    last_rl = 4 + len(STATES)
    for col, w in zip("ABCDEFGHIJ", [18, 12, 20, 14, 46, 18, 22, 34, 40, 46]):
        rl.column_dimensions[col].width = w
    rl.freeze_panes = "B5"

    wl["A1"] = "Direct Support Wages & Home Lease by State"; wl["A1"].font = Font(size=14, bold=True, color=NAVY)
    wl["A2"] = ("Wages: BLS OEWS May 2024 mean hourly wage, Home Health and Personal Care Aides (basis for DSP pay). "
                "Lease: estimated monthly lease for a 6-8 bed home (about 55% of the CLS 30-bed facility rent estimate). Both are editable on the P&L.")
    wl["A2"].font = F_NOTE
    for i, h in enumerate(["State", "Mean Hourly Wage", "Est. Monthly Home Lease (6-8 beds)"], 1):
        c = wl.cell(4, i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_S
    for r, st in enumerate(STATES, 5):
        wage, rent = W[st]
        wl.cell(r, 1, st); wl.cell(r, 2, wage).number_format = USD2
        wl.cell(r, 3, round(rent * 0.55 / 50) * 50).number_format = USD
    last_wl = 4 + len(STATES)
    wl.column_dimensions["A"].width = 20; wl.column_dimensions["B"].width = 16; wl.column_dimensions["C"].width = 32

    RATE_RNG = f"'ICF Rate Lookup'!$A$5:$J${last_rl}"
    WAGE_RNG = f"'Wage Lookup'!$A$5:$C${last_wl}"

    # ---------- P&L ----------
    ws.column_dimensions["A"].width = 44
    for c in ALLC:
        ws.column_dimensions[c].width = 11.5
    ws.column_dimensions["Q"].width = 50

    def band(r, text, fill=FILL_S, font=F_SEC):
        ws.merge_cells(f"A{r}:P{r}")
        c = ws[f"A{r}"]; c.value = text; c.fill = fill; c.font = font

    band(1, "Intermediate Care Facility (ICF/IID) - Pro Forma P&L", FILL_T, F_TITLE); ws.row_dimensions[1].height = 24
    band(2, "Small community ICF/IID group home | Medicaid per diem by state | Year 1 ramp-up + Years 2-3 | All figures are planning estimates", FILL_T, F_SUB)

    inputs = {}
    r = 4
    band(r, "SECTION A: STATE & HOME CONFIGURATION"); r += 1

    def inp(label, value, fmt=None, key=None, note=None, kind="in"):
        nonlocal r
        ws[f"A{r}"] = label; ws[f"A{r}"].font = F_N
        c = ws[f"B{r}"]; c.value = value; c.border = BOX
        if kind == "in":
            c.font = F_IN; c.fill = FILL_IN
        elif kind == "lk":
            c.font = F_LK
        else:
            c.font = F_B
        if fmt: c.number_format = fmt
        if note:
            ws[f"C{r}"] = note; ws[f"C{r}"].font = F_NOTE
        if key: inputs[key] = f"$B${r}"
        r += 1

    inp("State Selection", state_default, key="state", note="Dropdown - drives rate, wages and lease")
    dv = DataValidation(type="list", formula1=f"='ICF Rate Lookup'!$A$5:$A${last_rl}", allow_blank=False)
    ws.add_data_validation(dv); dv.add(inputs["state"])
    inp("Licensed Beds (number of residents at full occupancy)", 6, "0", "beds", "Small ICF/IID homes are typically 4-8 beds; up to 15")
    inp("State Medicaid Daily Rate (lookup)", f"=VLOOKUP({inputs['state']},{RATE_RNG},2,FALSE)", USD2, "lkrate", kind="lk")
    inp("Rate Status", f"=VLOOKUP({inputs['state']},{RATE_RNG},3,FALSE)", None, "status", kind="lk")
    inp("Rate Basis", f'=IF(VLOOKUP({inputs["state"]},{RATE_RNG},5,FALSE)="","See ICF Rate Lookup",VLOOKUP({inputs["state"]},{RATE_RNG},5,FALSE))', None, "basis", kind="lk")
    ws.merge_cells(f"B{r-1}:P{r-1}"); ws[f"B{r-1}"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r-1].height = 28
    inp("Small ICF Market Note", f'=IF(VLOOKUP({inputs["state"]},{RATE_RNG},7,FALSE)="","-",VLOOKUP({inputs["state"]},{RATE_RNG},7,FALSE))', None, "market", kind="lk")
    inp("Your Facility Daily Rate (override - leave blank to use lookup)", None, USD2, "override",
        "Enter your rate letter / level-of-need rate here once known")
    inp("Daily Rate Used", f"=IF(N({inputs['override']})>0,{inputs['override']},{inputs['lkrate']})", USD2, "rate", kind="calc")
    inp("Billable Days per Month", 30.42, "0.00", "days")
    inp("Annual Rate Increase (Years 2-3)", 0.025, PCT, "rinc")
    inp("Annual Cost Inflation (Years 2-3)", 0.03, PCT, "cinc")

    r += 1
    band(r, "SECTION B: OCCUPANCY RAMP-UP"); r += 1
    ws[f"A{r}"] = "";
    for i, c in enumerate(MC):
        ws[f"{c}{r}"] = f"Month {i+1}"
    ws[f"{Y1}{r}"] = "YEAR 1"; ws[f"{Y2}{r}"] = "YEAR 2"; ws[f"{Y3}{r}"] = "YEAR 3"
    for c in ["A"] + ALLC:
        ws[f"{c}{r}"].font = F_B; ws[f"{c}{r}"].fill = FILL_HD; ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    HDR_ROW = r; r += 1
    ramp = [0.34, 0.5, 0.5, 0.67, 0.67, 0.84, 0.84, 1, 1, 1, 1, 1]
    ws[f"A{r}"] = "Target Occupancy %"; OCC = r
    for c, v in zip(MC, ramp):
        ws[f"{c}{r}"] = v; ws[f"{c}{r}"].font = F_IN; ws[f"{c}{r}"].fill = FILL_IN; ws[f"{c}{r}"].number_format = "0%"
    ws[f"{Y1}{r}"] = f"=AVERAGE(B{r}:M{r})"; ws[f"{Y1}{r}"].number_format = PCT
    for c in (Y2, Y3):
        ws[f"{c}{r}"] = 0.95; ws[f"{c}{r}"].font = F_IN; ws[f"{c}{r}"].fill = FILL_IN; ws[f"{c}{r}"].number_format = "0%"
    ws[f"Q{r}"] = "Years 2-3: average occupancy (allows for turnover / hospital leave)"; ws[f"Q{r}"].font = F_NOTE
    r += 1
    ws[f"A{r}"] = "Residents Served"; RES = r
    for c in MC:
        ws[f"{c}{r}"] = f"=ROUND({inputs['beds']}*{c}{OCC},0)"; ws[f"{c}{r}"].number_format = "0"
    ws[f"{Y1}{r}"] = f"=AVERAGE(B{r}:M{r})"; ws[f"{Y1}{r}"].number_format = "0.0"
    for c in (Y2, Y3):
        ws[f"{c}{r}"] = f"={inputs['beds']}*{c}{OCC}"; ws[f"{c}{r}"].number_format = "0.0"
    r += 1
    ws[f"A{r}"] = "Resident Days"; DAYS = r
    for c in MC:
        ws[f"{c}{r}"] = f"={c}{RES}*{inputs['days']}"; ws[f"{c}{r}"].number_format = "#,##0"
    ws[f"{Y1}{r}"] = f"=SUM(B{r}:M{r})"; ws[f"{Y1}{r}"].number_format = "#,##0"
    for c in (Y2, Y3):
        ws[f"{c}{r}"] = f"={c}{RES}*365"; ws[f"{c}{r}"].number_format = "#,##0"
    r += 2

    band(r, "SECTION C: STAFFING & WAGE ASSUMPTIONS"); r += 1
    inp("DSP Hourly Wage (state lookup)", f"=VLOOKUP({inputs['state']},{WAGE_RNG},2,FALSE)", USD2, "wlk", kind="lk")
    inp("DSP Wage Adjustment (e.g. 1.10 = 10% above aide average)", 1.10, "0.00", "wadj", "ICF DSPs usually earn above the home care aide average")
    inp("DSP Hourly Wage Used", f"={inputs['wlk']}*{inputs['wadj']}", USD2, "dsp", kind="calc")
    inp("Awake Staff per Resident - Day & Evening Shifts (1 per N)", 3, "0", "rday", "42 CFR 483.430(d)(3): 1:3.2 to 1:6.4 by need level")
    inp("Awake Staff per Resident - Overnight Shift (1 per N)", 6, "0", "rnight", "Minimum one awake staff on every shift")
    inp("Day-Shift Hours Residents Spend at Day Program (%)", 0.35, PCT, "away",
        "Weekdays ~6 hrs off-site; in-home day-shift staffing is reduced by this share")
    inp("House Manager / Lead DSP Premium", 1.25, "0.00", "leadx")
    inp("QIDP (Qualified Intellectual Disability Professional) Salary", 62000, USD, "qidp", "Required by 42 CFR 483.430(a)")
    inp("QIDP FTE for this home", 0.5, "0.00", "qidpfte", "One QIDP often covers 2-3 homes")
    inp("Administrator / Program Director Salary", 72000, USD, "adm")
    inp("Administrator FTE for this home", 0.3, "0.00", "admfte")
    inp("Nurse (RN/LPN) Hourly Rate", f"={inputs['wlk']}*2.4", USD2, "rn", "Editable; about 2.4x aide wage")
    inp("Nurse Hours per Week", 12, "0", "rnhrs", "42 CFR 483.460 nursing services as needed")
    inp("Payroll Taxes & Benefits %", 0.20, PCT, "burden")
    r += 1

    def header(label):
        nonlocal r
        band(r, label); r += 1
        for i, c in enumerate(MC):
            ws[f"{c}{r}"] = f"Month {i+1}"
        ws[f"{Y1}{r}"] = "YEAR 1"; ws[f"{Y2}{r}"] = "YEAR 2"; ws[f"{Y3}{r}"] = "YEAR 3"
        for c in ["A"] + ALLC:
            ws[f"{c}{r}"].font = F_B; ws[f"{c}{r}"].fill = FILL_HD; ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
        r += 1

    def line(label, mformula, yformula, bold=False, fill=None, note=None, sum_y1=True):
        """mformula(c) -> monthly formula; yformula(c, k) -> year formula, k=1 for Y2, 2 for Y3"""
        nonlocal r
        ws[f"A{r}"] = label; ws[f"A{r}"].font = F_B if bold else F_N
        for c in MC:
            ws[f"{c}{r}"] = mformula(c)
        ws[f"{Y1}{r}"] = f"=SUM(B{r}:M{r})" if sum_y1 else mformula(Y1)
        ws[f"{Y2}{r}"] = yformula(Y2, 1); ws[f"{Y3}{r}"] = yformula(Y3, 2)
        for c in ALLC:
            ws[f"{c}{r}"].number_format = USD
            if bold: ws[f"{c}{r}"].font = F_B
            if fill: ws[f"{c}{r}"].fill = fill
        if fill: ws[f"A{r}"].fill = fill
        if note:
            ws[f"Q{r}"] = note; ws[f"Q{r}"].font = F_NOTE
        row = r; r += 1
        return row

    I = inputs
    rate = I["rate"]; ci = I["cinc"]; ri = I["rinc"]

    header("REVENUE")
    rev1 = line("Medicaid ICF/IID Per Diem Revenue", lambda c: f"={rate}*{c}{DAYS}",
                lambda c, k: f"={rate}*(1+{ri})^{k}*{c}{DAYS}")
    ws[f"{Y2}{OCC}"].number_format = "0%"
    I["oth"] = None
    ws[f"A{r}"] = "Other Revenue per Resident-Day (e.g. add-ons)"; ws[f"A{r}"].font = F_N
    # input stored in Q-adjacent? keep simple: input cell in B of assumptions block below; use 0 default via named row
    oth_in = r; r += 1
    ws.row_dimensions[oth_in].hidden = False
    for c in MC:
        ws[f"{c}{oth_in}"] = f"=$B${oth_in+40}*0"  # placeholder replaced below
    TOTREV = None

    # replace placeholder: other revenue uses an input cell in Section A-style block placed in column Q? Use simple design:
    # other revenue per resident-day input lives in column B of the 'Other inputs' row created next.
    r_other_input = r
    ws[f"A{r}"] = "  Input: Other revenue per resident-day ($)"; ws[f"A{r}"].font = F_NOTE
    ws[f"B{r}"] = 0; ws[f"B{r}"].font = F_IN; ws[f"B{r}"].fill = FILL_IN; ws[f"B{r}"].number_format = USD2; ws[f"B{r}"].border = BOX
    ws[f"C{r}"] = "Attendant-compensation enhancements, supplemental payments, etc."; ws[f"C{r}"].font = F_NOTE
    r += 1
    oth = f"$B${r_other_input}"
    for c in MC:
        ws[f"{c}{oth_in}"] = f"={oth}*{c}{DAYS}"
    ws[f"{Y1}{oth_in}"] = f"=SUM(B{oth_in}:M{oth_in})"
    ws[f"{Y2}{oth_in}"] = f"={oth}*(1+{ri})*{Y2}{DAYS}"; ws[f"{Y3}{oth_in}"] = f"={oth}*(1+{ri})^2*{Y3}{DAYS}"
    for c in ALLC:
        ws[f"{c}{oth_in}"].number_format = USD
    TOTREV = line("TOTAL REVENUE", lambda c: f"={c}{rev1}+{c}{oth_in}", lambda c, k: f"={c}{rev1}+{c}{oth_in}", True, FILL_TOT)
    r += 1

    header("STAFFING EXPENSES")
    H = f"8*{I['days']}"  # hours per shift-slot per month
    def staff_m(c, ratio):
        return f"=IF({c}{RES}=0,1,MAX(1,ROUNDUP({c}{RES}/{ratio},0)))"
    dsp_day = line("Day & Evening DSPs (2 shifts)",
                   lambda c: f"=(2-{I['away']})*{H}*{I['dsp']}*IF({c}{RES}=0,1,MAX(1,ROUNDUP({c}{RES}/{I['rday']},0)))",
                   lambda c, k: f"=(2-{I['away']})*8*365*{I['dsp']}*(1+{ci})^{k}*MAX(1,ROUNDUP({c}{RES}/{I['rday']},0))",
                   note="Staff on shift = residents / ratio, rounded up, minimum 1; day shift reduced while residents are at day program")
    dsp_ngt = line("Overnight Awake DSPs",
                   lambda c: f"={H}*{I['dsp']}*IF({c}{RES}=0,1,MAX(1,ROUNDUP({c}{RES}/{I['rnight']},0)))",
                   lambda c, k: f"=8*365*{I['dsp']}*(1+{ci})^{k}*MAX(1,ROUNDUP({c}{RES}/{I['rnight']},0))")
    lead = line("House Manager / Lead DSP (40 hrs/wk premium portion)",
                lambda c: f"=40*52/12*{I['dsp']}*({I['leadx']}-1)",
                lambda c, k: f"=40*52*{I['dsp']}*({I['leadx']}-1)*(1+{ci})^{k}",
                note="Lead works shifts already counted above; this is the premium only")
    qidp = line("QIDP", lambda c: f"={I['qidp']}*{I['qidpfte']}/12", lambda c, k: f"={I['qidp']}*{I['qidpfte']}*(1+{ci})^{k}")
    adm = line("Administrator / Program Director", lambda c: f"={I['adm']}*{I['admfte']}/12", lambda c, k: f"={I['adm']}*{I['admfte']}*(1+{ci})^{k}")
    nurse = line("Nursing (RN/LPN)", lambda c: f"={I['rn']}*{I['rnhrs']}*52/12", lambda c, k: f"={I['rn']}*{I['rnhrs']}*52*(1+{ci})^{k}")
    burden = line("Payroll Taxes & Benefits", lambda c: f"=SUM({c}{dsp_day}:{c}{nurse})*{I['burden']}",
                  lambda c, k: f"=SUM({c}{dsp_day}:{c}{nurse})*{I['burden']}")
    TOTSTAFF = line("TOTAL STAFFING", lambda c: f"=SUM({c}{dsp_day}:{c}{burden})", lambda c, k: f"=SUM({c}{dsp_day}:{c}{burden})", True, FILL_TOT)
    r += 1

    header("OPERATING EXPENSES")
    opex_rows = []
    lease_cell = f"VLOOKUP({I['state']},{WAGE_RNG},3,FALSE)"
    opex_rows.append(line("Home Lease / Mortgage (state lookup)", lambda c: f"={lease_cell}",
                          lambda c, k: f"={lease_cell}*12*(1+{ci})^{k}", note="6-8 bed home estimate; replace with your actual lease"))
    fixed = [
        ("Utilities (electric, gas, water, internet)", 900),
        ("Liability, Property & Auto Insurance", 1100),
        ("Vehicle (accessible van lease, fuel, upkeep)", 1300),
        ("Repairs, Maintenance & Life Safety Inspections", 500),
        ("Consultants (dietitian, pharmacist, psychologist)", 800),
        ("Staff Training & CPR / Medication Certification", 350),
        ("Background Checks & Drug Testing", 120),
        ("Software (EHR, eMAR, billing, scheduling)", 350),
        ("Licensing, Certification & Survey Costs", 150),
        ("Office, Phone & Administrative", 300),
        ("Accounting, Legal & Cost Report Preparation", 600),
        ("Miscellaneous / Contingency", 400),
    ]
    variable = [
        ("Food & Dietary (per resident-day)", 11.0),
        ("Household & Hygiene Supplies (per resident-day)", 3.0),
        ("Medical Supplies & OTC (per resident-day)", 2.5),
        ("Day Program / Active Treatment (per resident-day)", 30.0),
        ("Provider Fee / Bed Tax (per resident-day)", 0.0),
    ]
    # input block for opex lives in column Q? -> place inputs on an 'Expense Inputs' area below the P&L
    return_state = (wb, ws, rl, wl, ins, inputs, r, opex_rows, fixed, variable, ci, TOTREV, TOTSTAFF, DAYS, RES, HDR_ROW, line, band, header, path, last_rl, RATE_RNG)
    return finish(*return_state)


def finish(wb, ws, rl, wl, ins, I, r, opex_rows, fixed, variable, ci, TOTREV, TOTSTAFF, DAYS, RES, HDR_ROW, line, band, header, path, last_rl, RATE_RNG):
    # expense input table placed far below (rows 120+) so formulas can reference it
    base = 130
    ws[f"A{base-1}"] = "OPERATING EXPENSE INPUTS (edit the yellow cells)"; ws[f"A{base-1}"].font = Font(bold=True, color=NAVY, size=11)
    cells = {}
    rr = base
    for label, v in fixed:
        ws[f"A{rr}"] = label + " - monthly"; ws[f"A{rr}"].font = F_N
        ws[f"B{rr}"] = v; ws[f"B{rr}"].font = F_IN; ws[f"B{rr}"].fill = FILL_IN; ws[f"B{rr}"].number_format = USD; ws[f"B{rr}"].border = BOX
        cells[label] = f"$B${rr}"; rr += 1
    for label, v in variable:
        ws[f"A{rr}"] = label; ws[f"A{rr}"].font = F_N
        ws[f"B{rr}"] = v; ws[f"B{rr}"].font = F_IN; ws[f"B{rr}"].fill = FILL_IN; ws[f"B{rr}"].number_format = USD2; ws[f"B{rr}"].border = BOX
        cells[label] = f"$B${rr}"; rr += 1
    ws[f"C{base+len(fixed)+3}"] = "Set to $0 if your state's rate excludes day services and they are billed separately (e.g. IL, NY)"
    ws[f"C{base+len(fixed)+3}"].font = F_NOTE

    for label, _ in fixed:
        cell = cells[label]
        opex_rows.append(line(label, lambda c, cell=cell: f"={cell}", lambda c, k, cell=cell: f"={cell}*12*(1+{ci})^{k}"))
    for label, _ in variable:
        cell = cells[label]
        opex_rows.append(line(label, lambda c, cell=cell: f"={cell}*{c}{DAYS}", lambda c, k, cell=cell: f"={cell}*(1+{ci})^{k}*{c}{DAYS}"))
    o1, o2 = opex_rows[0], opex_rows[-1]
    TOTOPEX = line("TOTAL OPERATING EXPENSES", lambda c: f"=SUM({c}{o1}:{c}{o2})", lambda c, k: f"=SUM({c}{o1}:{c}{o2})", True, FILL_TOT)
    r_ = TOTOPEX + 2

    # profit summary
    import types
    def set_r(v):
        pass
    # continue writing with a local counter
    r = r_
    band(r, "PROFIT SUMMARY & PATH TO PROFITABILITY"); r += 1
    for i, c in enumerate(MC):
        ws[f"{c}{r}"] = f"Month {i+1}"
    ws[f"{Y1}{r}"] = "YEAR 1"; ws[f"{Y2}{r}"] = "YEAR 2"; ws[f"{Y3}{r}"] = "YEAR 3"
    for c in ["A"] + ALLC:
        ws[f"{c}{r}"].font = F_B; ws[f"{c}{r}"].fill = FILL_HD; ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    r += 1

    def row(label, f, fmt=USD, bold=False, fill=None, y1=None):
        nonlocal r
        ws[f"A{r}"] = label; ws[f"A{r}"].font = F_B if bold else F_N
        for c in ALLC:
            ws[f"{c}{r}"] = f(c)
            ws[f"{c}{r}"].number_format = fmt
            if bold: ws[f"{c}{r}"].font = F_B
            if fill: ws[f"{c}{r}"].fill = fill
        if y1: ws[f"{Y1}{r}"] = y1
        if fill: ws[f"A{r}"].fill = fill
        x = r; r += 1; return x

    TOTEXP = row("TOTAL EXPENSES", lambda c: f"={c}{TOTSTAFF}+{c}{TOTOPEX}", bold=True)
    NET = row("NET INCOME (LOSS)", lambda c: f"={c}{TOTREV}-{c}{TOTEXP}", bold=True, fill=FILL_TOT)
    CUM = r
    ws[f"A{r}"] = "CUMULATIVE P&L"; ws[f"A{r}"].font = F_B
    for i, c in enumerate(MC):
        ws[f"{c}{r}"] = f"={c}{NET}" if i == 0 else f"={MC[i-1]}{r}+{c}{NET}"
    ws[f"{Y1}{r}"] = f"=M{r}"; ws[f"{Y2}{r}"] = f"={Y1}{r}+{Y2}{NET}"; ws[f"{Y3}{r}"] = f"={Y2}{r}+{Y3}{NET}"
    for c in ALLC:
        ws[f"{c}{r}"].number_format = USD; ws[f"{c}{r}"].font = F_B
    r += 1
    MARGIN = row("Net Profit Margin", lambda c: f"=IF({c}{TOTREV}=0,0,{c}{NET}/{c}{TOTREV})", PCT)
    row("Staffing % of Revenue", lambda c: f"=IF({c}{TOTREV}=0,0,{c}{TOTSTAFF}/{c}{TOTREV})", PCT)
    row("Cost per Resident-Day", lambda c: f"=IF({c}{DAYS}=0,0,{c}{TOTEXP}/{c}{DAYS})", USD2)
    for rr_ in (NET, CUM):
        rng = f"B{rr_}:P{rr_}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=Font(color=GREEN, bold=True)))
    r += 1

    def kv(label, formula, fmt, note=""):
        nonlocal r
        ws[f"A{r}"] = label; ws[f"A{r}"].font = F_B
        ws[f"B{r}"] = formula; ws[f"B{r}"].number_format = fmt; ws[f"B{r}"].font = F_B; ws[f"B{r}"].border = BOX
        if note: ws[f"C{r}"] = note; ws[f"C{r}"].font = F_NOTE
        x = r; r += 1; return x

    kv("First Profitable Month", f'=IFERROR(MATCH(TRUE,INDEX(B{NET}:M{NET}>0,0),0),"After Month 12")', "0",
       "Month # when monthly net income first turns positive")
    kv("Year 2 Net Income (stabilized)", f"={Y2}{NET}", USD)
    ws[f"A{r}"] = "Reality Check"; ws[f"A{r}"].font = F_B
    ws[f"B{r}"] = (f'=IF({Y2}{MARGIN}>0.25,"Caution: most states set ICF/IID rates from each facility\'s allowable costs. A rate this far above '
                   f'your projected costs usually means higher-acuity residents and heavier required staffing, and cost-based rates are reset to '
                   f'actual costs after cost reports, so long-run margins are typically much thinner.",IF({Y2}{MARGIN}<0,"At this rate the home '
                   f'does not cover its costs. Check your level-of-need rate, add beds, or adjust staffing before proceeding.",'
                   f'"Margin is within a typical range for a small ICF/IID."))')
    ws.merge_cells(f"B{r}:P{r}"); ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws[f"B{r}"].font = Font(size=10, bold=True, color="9C5700"); ws.row_dimensions[r].height = 42
    r += 1
    kv("Year 3 Net Income", f"={Y3}{NET}", USD)
    # break-even residents at Year 2 economics: fixed + variable per resident
    occ_row = DAYS - 2
    kv("Break-Even Occupancy (Year 2 cost structure, approx.)",
       f"=IFERROR(MIN(1.5,{Y2}{TOTEXP}/{Y2}{TOTREV}*{Y2}{occ_row}),0)", PCT,
       "Approximate; staffing steps up in whole shifts, so check the monthly view")
    r += 1

    # startup costs
    band(r, "STARTUP COSTS & CASH NEEDED"); r += 1
    start = [
        ("Licensing, ICF/IID certification & consulting", 15000),
        ("Home modifications (accessibility, fire sprinkler / Life Safety Code)", 40000),
        ("Furniture, adaptive equipment & household setup", 18000),
        ("Lease deposit & first month", None),
        ("Accessible van down payment", 12000),
        ("Pre-opening hiring & training (about 1 month of core staff)", None),
        ("Policies, EHR/eMAR setup & insurance deposits", 6000),
    ]
    s0 = r
    for label, v in start:
        ws[f"A{r}"] = label; ws[f"A{r}"].font = F_N
        if v is None and label.startswith("Lease"):
            ws[f"B{r}"] = f"=2*{opex_rows[0] and 'B'+str(opex_rows[0])}"
            ws[f"B{r}"].font = F_N
        elif v is None:
            ws[f"B{r}"] = f"=B{TOTSTAFF}"
            ws[f"B{r}"].font = F_N
        else:
            ws[f"B{r}"] = v; ws[f"B{r}"].font = F_IN; ws[f"B{r}"].fill = FILL_IN
        ws[f"B{r}"].number_format = USD; ws[f"B{r}"].border = BOX
        r += 1
    START = kv("Total Startup Costs", f"=SUM(B{s0}:B{r-1})", USD)
    DEF = kv("Largest Cumulative Operating Deficit (Year 1)", f"=-MIN(0,MIN(B{CUM}:M{CUM}))", USD)
    RES_ = kv("Operating Reserve (3 months of Year 1 average expenses)", f"=3*{Y1}{TOTEXP}/12", USD,
              "Medicaid payment usually starts only after the certification survey; plan for a cash cushion")
    kv("TOTAL CASH NEEDED TO OPEN", f"=B{START}+MAX(B{DEF},B{RES_})", USD, "Startup + the larger of the deficit or the reserve")
    ws[f"B{r-1}"].fill = FILL_TOT
    r += 1
    ws[f"A{r}"] = ("Planning estimate only. Actual ICF/IID rates are set per facility by the state Medicaid agency (often from cost reports and "
                   "resident level of need). Confirm your rate, staffing requirements and certification timeline with your state before relying on these figures.")
    ws[f"A{r}"].font = F_NOTE; ws.merge_cells(f"A{r}:P{r}"); ws[f"A{r}"].alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 30

    ws.freeze_panes = ws[f"B{HDR_ROW+1}"]
    ws.sheet_view.zoomScale = 90

    # ---------- instructions ----------
    ins.column_dimensions["A"].width = 30; ins.column_dimensions["B"].width = 100
    lines = [
        ("ICF/IID P&L - User Guide", None),
        ("PURPOSE", "Projects Year 1 (monthly) and Years 2-3 for a small community ICF/IID group home paid by a Medicaid per diem."),
        ("QUICK START", None),
        ("1. Select State (B5)", "Dropdown sets the Medicaid daily rate, DSP wage and home lease from the lookup sheets."),
        ("2. Set Licensed Beds (B6)", "Number of residents at full occupancy. Everything scales from this: revenue, shifts staffed, food and supplies."),
        ("3. Enter your rate (B11)", "If you have a rate letter or a level-of-need rate, enter it. Otherwise the state lookup rate is used."),
        ("4. Edit the ramp (Section B)", "Monthly occupancy targets for Year 1; Years 2-3 average occupancy."),
        ("5. Review staffing (Section C)", "Shifts are staffed at residents / ratio, rounded up, minimum one awake staff per shift, 24/7."),
        ("6. Operating expenses", "Edit the yellow cells in the OPERATING EXPENSE INPUTS table at the bottom of the P&L."),
        ("RATE STATUS", None),
        ("Published", "Taken from a state rate schedule, rate letters or a state notice (see ICF Rate Lookup for the source and link)."),
        ("Proxy - low confidence", "Derived from state spending per resident. Use only until you have a real rate."),
        ("Statutory floor", "Minimum rate set in law; actual facility rates run higher."),
        ("Planning default - verify", "No usable published figure was found; a planning rate is used. Enter your facility's rate."),
        ("Small ICF market note", "Some states have no small private ICF/IID program (e.g. MI, OR, MA, NJ). Check before planning a new ICF there."),
        ("COST-BASED RATES", "Most states pay each ICF/IID a rate built from its own audited allowable costs (cost reports). High published rates "
         "usually reflect higher-acuity residents and heavier staffing, and rates are reset toward actual costs over time. Treat large projected "
         "margins with caution; flat-rate states (e.g. CA, TX, LA, OK) work differently."),
        ("KEY RULES (42 CFR 483 Subpart I)", "ICF/IIDs must provide 24-hour supervision, active treatment, a QIDP for each resident, nursing services as needed "
         "and direct-care staffing ratios of 1:3.2 to 1:6.4 by level of need. Most states also require a certificate of need or a state plan slot before certification."),
        ("COLOR CODING", "Blue text on yellow = editable input. Green text = looked up from reference sheets. Black = formula."),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ins[f"A{i}"] = a; ins[f"A{i}"].font = Font(bold=True, size=14 if i == 1 else 10, color=NAVY if (b is None) else "000000")
        if b: ins[f"B{i}"] = b; ins[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top"); ins[f"B{i}"].font = F_N

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def store_values(path):
    out = os.path.dirname(path) + "/_calc"
    os.makedirs(out, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, path],
                   check=True, capture_output=True, timeout=180)
    shutil.move(os.path.join(out, os.path.basename(path)), path)


if __name__ == "__main__":
    outdir = sys.argv[1]
    os.makedirs(outdir, exist_ok=True)
    master = os.path.join(outdir, "ICF_IID_PnL_All_States.xlsx")
    build("California", master); store_values(master)
    print("built", master)
    # one preset copy per state that sells the ICF bundle (slugs passed as remaining args)
    for slug in sys.argv[2:]:
        st_slug = slug.replace("-intermediate-care-facilities-complete-bundle", "")
        st = next(s for s in STATES if s.lower().replace(" ", "-") == st_slug)
        d = os.path.join(outdir, "bundles", slug); os.makedirs(d, exist_ok=True)
        p = os.path.join(d, f"{st} Intermediate Care Facilities - Pro Forma PnL.xlsx")
        build(st, p); store_values(p)
        print("built", p)
