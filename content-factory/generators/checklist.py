"""Licensing Checklist workbook (Oklahoma standard): phased Checklist with status tracking,
Rates Reference, Contacts and Sources."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from pnl import GREEN, HEAD, SECTION, USD2, f, header_row, put, src_ids, src_label

PRIORITY_FILL = {
    "CRITICAL": PatternFill("solid", fgColor="F8D7DA"),
    "HIGH": PatternFill("solid", fgColor="FFF3CD"),
    "MEDIUM": PatternFill("solid", fgColor="E2EFDA"),
}


def _phase_key(p: str) -> int:
    import re
    m = re.search(r"PHASE\s+(\d+)", p)
    return int(m.group(1)) if m else 999


def build_checklist(d: dict, path: str) -> None:
    sources = {s["id"]: s for s in d["sources"]}
    wb = Workbook()

    # ---- Checklist
    ws = wb.active
    ws.title = "Checklist"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (12, 52, 14, 13, 13, 85, 30, 26)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", f"{d['state']} {d['state_program_name']} — Provider Licensing Checklist", bold=True, size=14,
        color=GREEN)
    regs = ", ".join(r["cite"] for r in d.get("regulations", [])[:4])
    put(ws, "A2", f"Prepared by Care Licensing Solutions, {d.get('as_of', '')}. Governing rules: {regs}. "
                  "Set Status as you go; the counter at right tracks progress.", italic=True, wrap=True)
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 30
    header_row(ws, 4, ["Priority", "Task / Document", "Status", "Due Date", "Completed Date",
                       "Instructions / Notes", "Regulatory Citation", "Source"])
    ws.freeze_panes = "A5"

    items = sorted(enumerate(d["checklist"]), key=lambda t: (_phase_key(t[1]["phase"]), t[0]))
    r = 5
    phase = None
    task_rows = []
    for _, it in items:
        if it["phase"] != phase:
            phase = it["phase"]
            put(ws, f"A{r}", phase, bold=True, color="FFFFFF")
            for col in "ABCDEFGH":
                ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="2E6B4F")
            ws.merge_cells(f"A{r}:H{r}")
            r += 1
        pr = (it.get("priority") or "MEDIUM").upper()
        put(ws, f"A{r}", pr, bold=True, fill=PRIORITY_FILL.get(pr))
        put(ws, f"B{r}", it["task"], wrap=True, bold=pr == "CRITICAL")
        put(ws, f"C{r}", "Not started")
        put(ws, f"D{r}", None, fmt="mm/dd/yyyy")
        put(ws, f"E{r}", None, fmt="mm/dd/yyyy")
        put(ws, f"F{r}", it.get("instructions", ""), wrap=True)
        put(ws, f"G{r}", it.get("cite", ""), wrap=True)
        sid = it.get("source_id")
        c = put(ws, f"H{r}", src_label(sources, sid), italic=True, color="1155CC", wrap=True)
        ids = src_ids(sid)
        if ids and sources.get(ids[0], {}).get("url"):
            c.hyperlink = sources[ids[0]]["url"]
        ws.row_dimensions[r].height = max(30, min(150, 15 * (len(it.get("instructions", "")) // 95 + 1)))
        task_rows.append(r)
        r += 1

    dv = DataValidation(type="list", formula1='"Not started,In progress,Waiting on agency,Done,N/A"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    for tr in task_rows:
        dv.add(f"C{tr}")
    first, last = task_rows[0], task_rows[-1]
    ws.conditional_formatting.add(f"A{first}:H{last}", FormulaRule(
        formula=[f'$C{first}="Done"'], fill=PatternFill("solid", fgColor="D9D9D9"),
        font=f(color="808080")))
    ws.conditional_formatting.add(f"C{first}:C{last}", CellIsRule(
        operator="equal", formula=['"In progress"'], fill=PatternFill("solid", fgColor="FFF3CD")))
    # Progress counter
    put(ws, "G1", "Tasks done", bold=True)
    put(ws, "H1", f'=COUNTIF(C{first}:C{last},"Done")&" of "&(COUNTA(B{first}:B{last})'
                  f'-COUNTIF(C{first}:C{last},"N/A"))', bold=True, color=GREEN)
    put(ws, "G2", "Critical open", bold=True)
    put(ws, "H2", f'=COUNTIFS(A{first}:A{last},"CRITICAL",C{first}:C{last},"<>Done",'
                  f'C{first}:C{last},"<>N/A")', bold=True, color="C00000")
    ws.auto_filter.ref = f"A4:H{last}"

    # ---- Rates Reference
    rr = wb.create_sheet("Rates Reference")
    rr.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (46, 14, 14, 16, 14, 26, 50, 30)):
        rr.column_dimensions[col].width = w
    put(rr, "A1", f"{d['state']} {d['program']} — Published Rates", bold=True, size=14, color=GREEN)
    put(rr, "A2", "Confirm current rates with the payer before contracting; rates change with each fiscal year.",
        italic=True)
    header_row(rr, 4, ["Service", "Code", "Rate", "Unit", "Effective", "Payer", "Notes", "Source"])
    r = 5
    for rt in d.get("rates", []):
        put(rr, f"A{r}", rt.get("service", ""), wrap=True)
        put(rr, f"B{r}", rt.get("code", ""))
        put(rr, f"C{r}", rt.get("rate"), fmt=USD2)
        put(rr, f"D{r}", rt.get("unit", ""))
        put(rr, f"E{r}", rt.get("effective", ""))
        put(rr, f"F{r}", rt.get("payer", ""), wrap=True)
        put(rr, f"G{r}", rt.get("notes", ""), wrap=True)
        put(rr, f"H{r}", src_label(sources, rt.get("source_id")), italic=True, color="595959", wrap=True)
        r += 1
    rr.freeze_panes = "A5"

    # ---- Contacts
    co = wb.create_sheet("Contacts")
    co.sheet_view.showGridLines = False
    for col, w in zip("ABC", (42, 60, 60)):
        co.column_dimensions[col].width = w
    put(co, "A1", "Key Contacts & Authoritative Resources", bold=True, size=14, color=GREEN)
    header_row(co, 3, ["Entity / Resource", "Contact", "Role / Notes"])
    r = 4
    put(co, f"A{r}", "SECTION: REGULATORS", bold=True, fill=SECTION)
    co.merge_cells(f"A{r}:C{r}")
    r += 1
    for g in d.get("regulators", []):
        contact = " / ".join(x for x in [g.get("phone"), g.get("email"), g.get("url")] if x)
        put(co, f"A{r}", g["name"], wrap=True, bold=True)
        put(co, f"B{r}", contact, wrap=True)
        put(co, f"C{r}", g.get("role", ""), wrap=True)
        r += 1
    if d.get("contacts"):
        put(co, f"A{r}", "SECTION: OTHER CONTACTS", bold=True, fill=SECTION)
        co.merge_cells(f"A{r}:C{r}")
        r += 1
        for c in d["contacts"]:
            put(co, f"A{r}", c["name"], wrap=True, bold=True)
            put(co, f"B{r}", c.get("contact", ""), wrap=True)
            put(co, f"C{r}", c.get("role", ""), wrap=True)
            r += 1
    if d.get("regulations"):
        r += 1
        put(co, f"A{r}", "SECTION: GOVERNING RULES", bold=True, fill=SECTION)
        co.merge_cells(f"A{r}:C{r}")
        r += 1
        for g in d["regulations"]:
            put(co, f"A{r}", g["cite"], bold=True, wrap=True)
            c = put(co, f"B{r}", g.get("url", ""), color="1155CC", wrap=True)
            if g.get("url"):
                c.hyperlink = g["url"]
            put(co, f"C{r}", g.get("title", ""), wrap=True)
            r += 1

    # ---- Sources
    so = wb.create_sheet("Sources")
    so.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (7, 60, 34, 70, 12)):
        so.column_dimensions[col].width = w
    put(so, "A1", "Sources", bold=True, size=14, color=GREEN)
    header_row(so, 3, ["ID", "Title", "Publisher", "URL", "Accessed"])
    for i, s in enumerate(d["sources"]):
        row = 4 + i
        put(so, f"A{row}", s["id"])
        put(so, f"B{row}", s.get("title", ""), wrap=True)
        put(so, f"C{row}", s.get("publisher", ""), wrap=True)
        c = put(so, f"D{row}", s.get("url", ""), color="1155CC")
        if s.get("url"):
            c.hyperlink = s["url"]
        put(so, f"E{row}", s.get("accessed", ""))

    wb.save(path)
